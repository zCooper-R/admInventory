import io

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.import_export.models import ImportLog, ImportStatus
from apps.import_export.services import process_excel_import
from apps.inventory.models import Browser, Device, InternetSpeed, Position
from apps.inventory.services.export import export_pcs_to_excel
from apps.inventory.services.template_excel import COLUMNS, NUMBERING_ROW
from .factories import AdminUserFactory, DeviceFactory


@pytest.mark.django_db
class TestExportMatchesTemplate:
    def test_export_has_same_header_and_numbering_row_as_template(self):
        DeviceFactory(inventory_number="EXP-001")

        payload = export_pcs_to_excel()
        wb = openpyxl.load_workbook(io.BytesIO(payload))
        ws = wb.active

        header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        numbering = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]

        assert header == COLUMNS
        assert numbering == NUMBERING_ROW

    def test_export_file_can_be_imported_back(self):
        browser = Browser.objects.create(name="Яндекс Браузер", normalized_name="яндекс браузер")
        position = Position.objects.create(name="Инженер", normalized_name="инженер")

        source = DeviceFactory(
            inventory_number="EXP-ROUNDTRIP-001",
            os="Windows 11 Pro",
            cpu_model="Intel Core i5-12400",
            cpu_frequency=2.50,
            ram=16,
            storage_type="SSD",
            storage_size=512,
            browser=browser,
            position=position,
            internet_speed=InternetSpeed.ABOVE_100,
            provider="Ростелеком",
            employee_name="Иванов И.И.",
            has_google_account=True,
            has_apple_account=False,
            has_microsoft_account=True,
            is_certified=True,
            use_for_text=True,
            use_for_images=False,
            use_for_presentations=True,
            use_for_audio=False,
            use_for_video=False,
        )

        exported_bytes = export_pcs_to_excel()

        user = AdminUserFactory()
        uploaded = SimpleUploadedFile(
            "roundtrip.xlsx",
            exported_bytes,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        log = ImportLog.objects.create(file=uploaded, uploaded_by=user)

        process_excel_import(log)
        log.refresh_from_db()

        assert log.status in {ImportStatus.SUCCESS, ImportStatus.PARTIAL}
        assert log.error_count == 0

        reloaded = Device.objects.get(inventory_number="EXP-ROUNDTRIP-001")
        assert reloaded.organization_id == source.organization_id
        assert reloaded.os == "Windows 11 Pro"
        assert reloaded.cpu_model.startswith("Intel Core i5")
        assert reloaded.storage_type == "SSD"
        assert reloaded.browser is not None
        assert reloaded.position is not None
