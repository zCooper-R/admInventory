"""
Tests for the Excel import template download endpoint and service.

Covers
------
* Unauthenticated request is redirected.
* Response has the correct Content-Type and Content-Disposition.
* ``build_import_template_bytes`` returns valid xlsx bytes with expected columns.

Author : Литвин Олег Олегович <qucooper@yandex.ru>
"""

import io

import openpyxl
import pytest
from apps.inventory.services.template_excel import COLUMNS, build_import_template_bytes
from django.test import Client
from django.urls import reverse

from .factories import AdminUserFactory

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def auth_client(db):
    user = AdminUserFactory()
    c = Client()
    c.force_login(user)
    return c


@pytest.mark.django_db
class TestImportTemplateView:
    def test_requires_login(self):
        url = reverse("pc-import-template")
        response = Client().get(url)
        assert response.status_code == 302
        assert "/login/" in response["Location"]

    def test_returns_200_for_authenticated(self, auth_client):
        url = reverse("pc-import-template")
        response = auth_client.get(url)
        assert response.status_code == 200

    def test_content_type_is_xlsx(self, auth_client):
        url = reverse("pc-import-template")
        response = auth_client.get(url)
        assert XLSX_CONTENT_TYPE in response["Content-Type"]

    def test_content_disposition_is_attachment(self, auth_client):
        url = reverse("pc-import-template")
        response = auth_client.get(url)
        disposition = response.get("Content-Disposition", "")
        assert "attachment" in disposition
        assert ".xlsx" in disposition

    def test_response_is_valid_xlsx(self, auth_client):
        url = reverse("pc-import-template")
        response = auth_client.get(url)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        assert wb is not None

    def test_first_row_contains_expected_columns(self, auth_client):
        url = reverse("pc-import-template")
        response = auth_client.get(url)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        header_row = [cell.value for cell in next(ws.iter_rows(max_row=1))]
        for col in COLUMNS:
            assert col in header_row, f"Column '{col}' missing from template header"


class TestBuildImportTemplateBytes:
    """Unit tests for the service function — no DB required."""

    def test_returns_bytes(self):
        result = build_import_template_bytes()
        assert isinstance(result, bytes)

    def test_non_empty(self):
        result = build_import_template_bytes()
        assert len(result) > 0

    def test_produces_valid_xlsx(self):
        result = build_import_template_bytes()
        wb = openpyxl.load_workbook(io.BytesIO(result))
        assert wb is not None

    def test_has_sample_data(self):
        result = build_import_template_bytes()
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) >= 1
        assert any(any(cell for cell in row) for row in rows)
