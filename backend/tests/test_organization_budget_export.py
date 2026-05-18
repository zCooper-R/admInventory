import io

import openpyxl
import pytest
from apps.inventory.models import Device, ReplacementStatus, SystemSettings
from apps.inventory.services.template_excel import COLUMNS
from django.test import Client
from django.urls import reverse

from .factories import AdminUserFactory, DeviceFactory, OrganizationFactory

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


@pytest.fixture
def auth_client(db):
    client = Client()
    client.force_login(AdminUserFactory())
    return client


def _workbook(response):
    return openpyxl.load_workbook(io.BytesIO(response.content))


@pytest.mark.django_db
class TestOrganizationBudgetExport:
    def test_old_export_endpoint_still_returns_xlsx(self, auth_client):
        DeviceFactory()

        response = auth_client.get(reverse("pc-export"))

        assert response.status_code == 200
        assert response["Content-Type"] == XLSX_CONTENT_TYPE
        assert response["Content-Disposition"].endswith('.xlsx"')

    def test_old_export_endpoint_keeps_original_columns(self, auth_client):
        DeviceFactory()

        response = auth_client.get(reverse("pc-export"))
        ws = _workbook(response).active

        assert [cell.value for cell in ws[1]] == COLUMNS

    def test_new_endpoint_returns_xlsx_with_expected_columns(self, auth_client):
        DeviceFactory()

        response = auth_client.get(reverse("organization-budget-export"))
        ws = _workbook(response).active

        assert response.status_code == 200
        assert response["Content-Type"] == XLSX_CONTENT_TYPE
        assert [cell.value for cell in ws[5]] == [
            "Организация",
            "Общее кол-во ус-в",
            "Плохих",
            "Сколько денег нужно",
        ]

    def test_bad_devices_include_only_replace_status(self, auth_client):
        org = OrganizationFactory(name="Организация отчёта")
        replace_device = DeviceFactory(organization=org)
        attention_device = DeviceFactory(organization=org)
        ok_device = DeviceFactory(organization=org)
        Device.objects.filter(pk=replace_device.pk).update(
            replacement_status=ReplacementStatus.REPLACE
        )
        Device.objects.filter(pk=attention_device.pk).update(
            replacement_status=ReplacementStatus.ATTENTION
        )
        Device.objects.filter(pk=ok_device.pk).update(
            replacement_status=ReplacementStatus.OK
        )

        response = auth_client.get(
            reverse("organization-budget-export"), {"price": "12345"}
        )
        ws = _workbook(response).active

        data_row = [cell.value for cell in ws[6]]
        assert data_row == ["Организация отчёта", 3, 1, 12345]

    def test_get_price_changes_required_money(self, auth_client):
        org = OrganizationFactory(name="Цена из GET")
        d1 = DeviceFactory(organization=org)
        d2 = DeviceFactory(organization=org)
        Device.objects.filter(pk__in=[d1.pk, d2.pk]).update(
            replacement_status=ReplacementStatus.REPLACE
        )

        response = auth_client.get(
            reverse("organization-budget-export"), {"price": "70000"}
        )
        ws = _workbook(response).active

        assert ws["D6"].value == 140000

    def test_default_price_uses_system_settings(self, auth_client):
        SystemSettings.objects.update_or_create(pk=1, defaults={"pc_price_default": 34567})
        org = OrganizationFactory(name="Цена из настроек")
        device = DeviceFactory(organization=org)
        Device.objects.filter(pk=device.pk).update(
            replacement_status=ReplacementStatus.REPLACE
        )

        response = auth_client.get(reverse("organization-budget-export"))
        ws = _workbook(response).active

        assert ws["D6"].value == 34567

    def test_budget_page_has_generate_report_button(self, auth_client):
        response = auth_client.get(reverse("budget-report"))
        body = response.content.decode("utf-8")

        assert response.status_code == 200
        assert 'id="organization-report-export-btn"' in body
        assert "Сформировать отчёт" in body
        assert reverse("organization-budget-export") in body
