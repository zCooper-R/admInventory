from __future__ import annotations

import logging
from typing import Any

from openpyxl import load_workbook

from .types import ParsedRow, ParseResult

logger = logging.getLogger(__name__)

HEADER_ALIASES: dict[str, str] = {
    "наименование юридического лица": "organization_name",
    "адрес организации": "organization_address",
    "адрес": "organization_address",
    "площадка": "organization_address",
    "инв. №": "inventory_number",
    "инв №": "inventory_number",
    "инвентарный номер": "inventory_number",
    "наименование ос": "os",
    "наименование процессора": "cpu_model",
    "тактовая частота": "cpu_frequency",
    "оперативная память": "ram",
    "тип диска": "storage_type",
    "емкость диска": "storage_size",
    "ёмкость диска": "storage_size",
    "браузер которым пользуетесь": "browser",
    "наличие личного аккаунта google, аккаунта apple или аккаунта microsoft": "accounts",
    "скорость интернета": "internet_speed",
    "провайдер": "provider",
    "аттестованный компьютер": "is_certified",
    "работа с текстом": "use_for_text",
    "работа с картинками, фотографиями": "use_for_images",
    "создание презентаций": "use_for_presentations",
    "работа с аудио": "use_for_audio",
    "работа с видео": "use_for_video",
    "фамилия, инициалы сотрудника": "employee_name",
    "должность": "position",
}

REQUIRED_HEADERS = {
    "organization_name",
    "inventory_number",
    "os",
    "ram",
    "storage_type",
}
DEVICE_ROW_KEYS = {
    "organization_name",
    "inventory_number",
    "os",
    "ram",
    "storage_type",
    "employee_name",
    "position",
}


def _normalize_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("ё", "е").split())


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_numbering_row(values: list[str]) -> bool:
    filtered = [v for v in values if v]
    if not filtered:
        return False
    numeric = [v for v in filtered if v.isdigit()]
    return len(filtered) >= 5 and len(numeric) >= max(5, len(filtered) - 1)


def _is_device_row(values: dict[str, str]) -> bool:
    return any(_normalize_cell(values.get(key, "")) for key in DEVICE_ROW_KEYS)


def parse_excel(path: str) -> ParseResult:
    logger.info("Начало парсинга Excel: path=%s", path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return ParseResult(rows=[], total_rows=0)

    headers = [_normalize_header(v) for v in header_row]
    header_idx: dict[str, int] = {}
    for idx, raw_header in enumerate(headers):
        if not raw_header:
            continue
        mapped = HEADER_ALIASES.get(raw_header)
        if mapped:
            header_idx[mapped] = idx

    missing = REQUIRED_HEADERS - set(header_idx)
    if missing:
        missing_cols = ", ".join(sorted(missing))
        raise ValueError(
            f"Отсутствуют обязательные колонки нового формата: {missing_cols}"
        )

    parsed_rows: list[ParsedRow] = []
    total_rows = 0
    skipped_empty_rows = 0
    skipped_numbering_rows = 0
    skipped_non_device_rows = 0

    for excel_row_number, raw_row in enumerate(rows_iter, start=2):
        row_values = [_normalize_cell(v) for v in raw_row]
        row_data = {
            field_name: row_values[idx] if idx < len(row_values) else ""
            for field_name, idx in header_idx.items()
        }

        mapped_values = [
            row_data.get(field_name, "") for field_name in header_idx.keys()
        ]
        if not any(mapped_values):
            skipped_empty_rows += 1
            continue
        if _is_numbering_row(mapped_values):
            skipped_numbering_rows += 1
            continue
        if not _is_device_row(row_data):
            skipped_non_device_rows += 1
            continue

        total_rows += 1
        parsed_rows.append(ParsedRow(row_number=excel_row_number, values=row_data))

    logger.info(
        (
            "Парсинг завершён: total_rows=%s, parsed=%s, skipped_empty=%s, "
            "skipped_numbering=%s, skipped_non_device=%s"
        ),
        total_rows,
        len(parsed_rows),
        skipped_empty_rows,
        skipped_numbering_rows,
        skipped_non_device_rows,
    )
    return ParseResult(
        rows=parsed_rows,
        total_rows=total_rows,
        skipped_empty_rows=skipped_empty_rows,
        skipped_numbering_rows=skipped_numbering_rows,
        skipped_non_device_rows=skipped_non_device_rows,
    )
