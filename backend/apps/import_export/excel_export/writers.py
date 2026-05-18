from __future__ import annotations

import io
from collections.abc import Iterable

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


def workbook_to_bytes(workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.read()


def autosize_columns(ws: Worksheet, columns: Iterable[int] | None = None) -> None:
    column_indexes = columns or range(1, ws.max_column + 1)
    for column_idx in column_indexes:
        letter = get_column_letter(column_idx)
        max_len = 0
        for cell in ws.iter_cols(
            min_col=column_idx, max_col=column_idx, min_row=1, max_row=ws.max_row
        ):
            for item in cell:
                value = "" if item.value is None else str(item.value)
                max_len = max(max_len, len(value))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 60)
