from __future__ import annotations

from collections.abc import Sequence
from datetime import datetime

from apps.inventory.services.budget_report import OrganizationBudgetReportRow
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .writers import autosize_columns, workbook_to_bytes

HEADERS = [
    "Организация",
    "Общее кол-во ус-в",
    "Плохих",
    "Сколько денег нужно",
]


def export_organization_budget_to_excel(
    rows: Sequence[OrganizationBudgetReportRow],
    price: int,
    generated_at: datetime | None = None,
) -> bytes:
    generated_at = generated_at or datetime.now()

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"

    ws["A1"] = "Отчёт по бюджету замен"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")

    ws["A2"] = f"Дата формирования: {generated_at:%d.%m.%Y %H:%M}"
    ws.merge_cells("A2:D2")
    ws["A3"] = f"Цена за единицу: {price} ₽"
    ws.merge_cells("A3:D3")

    header_row = 5
    ws.append([])
    ws.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="E2E8F0")
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    total_devices = 0
    total_bad = 0
    total_money = 0

    for row in rows:
        total_devices += row.total_devices
        total_bad += row.bad_devices
        total_money += row.required_money
        ws.append(
            [
                row.organization_name,
                row.total_devices,
                row.bad_devices,
                row.required_money,
            ]
        )

    ws.append(["Итого", total_devices, total_bad, total_money])
    total_row = ws.max_row
    for cell in ws[total_row]:
        cell.font = Font(bold=True)

    money_column = 4
    for row_idx in range(header_row + 1, ws.max_row + 1):
        ws.cell(row=row_idx, column=money_column).number_format = '#,##0" ₽"'

    for row in ws.iter_rows(
        min_row=header_row, max_row=ws.max_row, min_col=2, max_col=4
    ):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")

    ws.freeze_panes = "A6"
    autosize_columns(ws)

    return workbook_to_bytes(wb)
